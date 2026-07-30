from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "product" / "launch-plan.md"
OUTPUT = ROOT / "product" / "StressFreeClaim-Production-Launch-Plan.docx"

OXBLOOD = RGBColor(0x5E, 0x2A, 0x28)
CORAL = RGBColor(0xE0, 0x71, 0x4C)
TERRACOTTA = RGBColor(0xB8, 0x4E, 0x2E)
INK = RGBColor(0x2E, 0x23, 0x20)
MUTED = RGBColor(0x7C, 0x6B, 0x62)
BLUSH = "FBE7DC"
LIGHT = "FDF7F2"
BORDER = "EFD9CC"
WHITE = "FFFFFF"


def set_font(run, name="Aptos", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("StressFreeClaim.ai  •  Production launch plan")
    set_font(run, size=8.5, color=MUTED)


def add_header(section):
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("DESKAR 6 LLC  /  CONFIDENTIAL WORKING PLAN")
    set_font(run, size=8.5, color=MUTED, bold=True)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_header(section)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (
        ("Heading 1", 16, 14, 6),
        ("Heading 2", 12.5, 10, 4),
        ("Heading 3", 11, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Georgia"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = OXBLOOD
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.25)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("StressFreeClaim.ai")
    set_font(r, "Georgia", 28, OXBLOOD, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Production Launch Plan")
    set_font(r, "Georgia", 18, OXBLOOD)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Prepared for Deskar 6 LLC  •  Recommended scope for approval  •  July 28, 2026")
    set_font(r, size=9.5, color=MUTED, bold=True)

    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2100, 7260])
    shade(table.cell(0, 0), "5E2A28")
    shade(table.cell(0, 1), BLUSH)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    r = left.add_run("RECOMMENDATION")
    set_font(r, size=9, color=WHITE and RGBColor(255, 255, 255), bold=True)
    right = table.cell(0, 1).paragraphs[0]
    right.paragraph_format.space_after = Pt(0)
    r = right.add_run(
        "Launch a counsel-cleared private pilot with durable intake, authenticated staff access, "
        "and an operational claims queue. Keep contracts, fees, insurer filing, voice, and CRM out."
    )
    set_font(r, size=10.25, color=INK, bold=True)


def add_callout(doc, label, body, fill=LIGHT):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1700, 7660])
    shade(table.cell(0, 0), "B84E2E")
    shade(table.cell(0, 1), fill)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label.upper())
    set_font(r, size=8.5, color=RGBColor(255, 255, 255), bold=True)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(body)
    set_font(r, size=10, color=INK)


def add_phase_table(doc):
    rows = [
        ("0", "Paper & operations", "2–5 business days, parallel", "States, copy, privacy, owners, signed SOW"),
        ("1", "Production foundation", "3–5 engineering days", "Durable data, auth, audit, environments"),
        ("2", "Launch slice", "5–8 engineering days", "Intake, eligibility, queue, workflow, notifications"),
        ("3", "Hardening & release", "3–5 engineering days", "Security, restore, QA, training, runbooks"),
        ("Pilot", "Controlled private pilot", "2 weeks", "Manual review, metrics, twice-weekly decisions"),
    ]
    table = doc.add_table(rows=1, cols=4)
    widths = [760, 2380, 2140, 4080]
    set_table_geometry(table, widths)
    headers = ["Stage", "Work", "Target", "Exit signal"]
    for i, text in enumerate(headers):
        shade(table.cell(0, i), "5E2A28")
        p = table.cell(0, i).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=9, color=RGBColor(255, 255, 255), bold=True)
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                shade(cells[i], LIGHT)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_font(r, size=9.2, bold=(i == 0))


def add_bullet(doc, text, numbered=False, compact=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.keep_together = True
    if compact:
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_font(r, size=9.5 if compact else 10.25)
    return p


def build():
    doc = Document()
    style_document(doc)
    add_title_block(doc)

    doc.add_heading("Executive decision", level=1)
    doc.add_paragraph(
        "Phase 1 is complete. The next release should be a narrow production-grade intake and "
        "operations product—not a broader automation platform. The fastest responsible route to "
        "market is a private pilot in counsel-cleared states."
    )
    for text in [
        "Preserve the proven homeowner utterance → gaps → review experience.",
        "Add durable storage, authenticated staff access, a claims queue, notifications, auditability, and production operations.",
        "Exclude contracts, fee language, insurer filing, voice, CRM, and partner tools from launch.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("Why the original plan changes", level=1)
    doc.add_paragraph(
        "The original Phase 2 correctly called for persistent claims and an internal view, but its "
        "one-to-two-week estimate describes visible features rather than everything required to "
        "accept real homeowner data. Today, drafts contain personal information in a browser cookie, "
        "submitted claims use ephemeral SQLite storage, there is no authentication, and the repository "
        "constitution intentionally prohibits production data and security work."
    )
    add_callout(
        doc,
        "Promotion rule",
        "Carry the validated intake UX forward. Replace the prototype storage, identity, security, "
        "and operations substrate under a separate production track and acceptance gate.",
    )

    doc.add_heading("Launch promise", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        "“Tell us what happened. A member of our team will review your information and contact you about next steps.”"
    )
    set_font(r, "Georgia", 12, OXBLOOD, italic=True)
    doc.add_paragraph(
        "Until counsel and operations approve otherwise, do not promise that a claim has been filed, "
        "that the company will represent the homeowner, that an affiliate will perform repairs, or "
        "that a fee will be waived."
    )

    doc.add_heading("Launch scope", level=1)
    doc.add_heading("Homeowner experience", level=2)
    for text in [
        "Mobile-first guided intake with a server-side, expiring draft; the browser cookie carries only a random continuation token.",
        "State eligibility against an administratively controlled allowlist.",
        "Clear consent and privacy disclosure, full review and correction, and a truthful submission receipt.",
        "Email confirmation plus calm recovery for expired drafts, duplicates, service failures, and unsupported states.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Staff experience", level=2)
    for text in [
        "Authenticated staff access with approved users, server-side roles, and MFA support.",
        "Claims queue with status, state, claimant, date of loss, insurer, submitted time, assignment, search, and filters.",
        "Claim detail, internal notes, assignment, and status workflow: new, reviewing, contacted, qualified, closed, duplicate.",
        "Audit-recorded access, material changes, and authorized CSV exports.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Platform and operations", level=2)
    for text in [
        "Managed PostgreSQL, migrations, backups, and separate development, preview, and production environments.",
        "Transactional email, structured logs, error monitoring, uptime checks, and personal-data redaction in telemetry.",
        "Rate limiting, secure headers, request protections, server-side validation, retention/deletion procedures, and environment-managed secrets.",
        "A named operator reviews the pilot queue every business day.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("Explicitly out of scope", level=1)
    for text in [
        "Public-adjuster or construction contracts, e-signatures, and cancellation-window workflows.",
        "Fee calculations, fee waivers, referral economics, or affiliate financial language.",
        "Automated insurer calls or electronic claim filing; voice capture, transcription, or recording.",
        "Payments, contractor portal, CRM sync, native apps, homeowner accounts, and multi-tenant organizations.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("Recommended delivery sequence", level=1)
    add_phase_table(doc)
    doc.add_paragraph(
        "Recommended target: roughly three engineering weeks to a controlled pilot, subject to the "
        "legal, privacy, operational, and release gates—not a calendar-only promise."
    )

    doc.add_heading("Workstream outcomes", level=1)
    outcomes = [
        ("Gate 0 — Paper & operations", "Pilot states, approved copy, receiving entity, privacy and retention decisions, staff list, queue owner, support promise, incident contacts, signed MSA/SOW, and client-owned production accounts."),
        ("Workstream 1 — Foundation", "A staff-only production shell authenticates correctly, stores synthetic claims durably, separates environments, and emits traceable audit events."),
        ("Workstream 2 — Launch slice", "The full mobile intake through staff follow-up works with synthetic data, including eligibility, consent, idempotency, status, assignment, notes, and notifications."),
        ("Workstream 3 — Hardening", "Authorization, accessibility, cross-browser and failure-mode QA, backup restore, telemetry redaction, incident rehearsal, staff training, and end-to-end launch rehearsal all pass."),
        ("Private pilot", "Controlled traffic in approved states, manual review of every submission, and twice-weekly review of completion, contact time, qualification, duplicates, support incidents, and abandonment."),
    ]
    for title, body in outcomes:
        doc.add_heading(title, level=2)
        doc.add_paragraph(body)

    doc.add_heading("Launch acceptance gates", level=1)
    gates = [
        "A submitted claim survives deployment, restart, and a successful database restore rehearsal.",
        "Personal information is absent from the browser draft cookie and client-visible logs.",
        "Anonymous users cannot read claims; unauthenticated staff cannot access protected surfaces.",
        "Staff roles are enforced server-side for every protected read and write.",
        "Every persisted homeowner field is visible and correctable before submission.",
        "Duplicate submissions create one claim and one notification set.",
        "Unsupported states cannot submit a production claim.",
        "Notification failures are retried safely and visible to operators.",
        "Claim access, exports, assignments, status changes, and notes generate audit events.",
        "Backups, privacy, retention, deletion, monitoring, and incident response have named owners and rehearsed procedures.",
        "Production copy contains no unapproved filing, representation, contractor, or fee promise.",
        "The prototype remains a separate demo; production carries a pilot label and support contact.",
    ]
    for gate in gates:
        add_bullet(doc, gate, numbered=True, compact=True)

    doc.add_heading("Suggested production architecture", level=1)
    architecture = [
        ("Application", "Keep Next.js, TypeScript, Tailwind, and Vercel."),
        ("Data", "Use managed PostgreSQL and retain Prisma initially to minimize migration risk."),
        ("Identity", "Use a managed provider with server-side sessions, staff allowlisting, roles, and MFA support."),
        ("Email", "Use a transactional provider with delivery events and idempotent sending."),
        ("Monitoring", "Use error monitoring and structured logs with deliberate personal-data redaction."),
        ("Extraction", "Keep deterministic local parsing authoritative. The bounded model pass stays optional and cannot block or overwrite local values."),
    ]
    for label, value in architecture:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{label}: ")
        set_font(r, size=9.5, bold=True, color=OXBLOOD)
        r = p.add_run(value)
        set_font(r, size=9.5)

    doc.add_heading("Decisions needed from Deskar 6 LLC", level=1)
    decisions = [
        "Pilot states and the counsel who approved them.",
        "Exact public-facing entity and brand.",
        "Named staff users and the pilot queue owner.",
        "Response-time promise to homeowners.",
        "Privacy/retention owner and deletion contact.",
        "Whether policy number and deductible should remain optional or be collected at all.",
        "Desired pilot start date and expected weekly claim volume.",
    ]
    for decision in decisions:
        add_bullet(doc, decision, compact=True)

    doc.core_properties.title = "StressFreeClaim.ai Production Launch Plan"
    doc.core_properties.subject = "Recommended MVP promotion and private-pilot plan"
    doc.core_properties.author = "Prepared for Deskar 6 LLC"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
